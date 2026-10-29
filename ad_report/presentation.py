"""Подстановка данных в шаблон презентации (только существующие элементы)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData

from .audience import Audience, MonthRange
from .metrics import Report


def _set_text(paragraph, text: str) -> None:
    """Меняет текст абзаца, сохраняя форматирование первого run."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        paragraph.add_run().text = text


def build_presentation(template: Path, out: Path, audience: Audience,
                       period: MonthRange, report: Report) -> None:
    prs = Presentation(str(template))

    subtitle = prs.slides[0].shapes[1].text_frame
    _set_text(subtitle.paragraphs[0], f"Целевая аудитория: {audience.describe()}")
    _set_text(subtitle.paragraphs[1], f"Период: {period.describe()}")

    table = prs.slides[1].shapes[1].table
    names = list(report.top_reach.index)
    for i in range(1, len(table.rows)):
        cells = table.rows[i].cells
        if i - 1 < len(names):
            _set_text(cells[0].text_frame.paragraphs[0], names[i - 1])
            _set_text(cells[1].text_frame.paragraphs[0], f"{report.top_reach.iloc[i - 1] * 100:.2f}%")
        else:  # рекламодателей меньше пяти — лишние строки очищаем
            _set_text(cells[0].text_frame.paragraphs[0], "")
            _set_text(cells[1].text_frame.paragraphs[0], "")

    chart = prs.slides[2].shapes[1].chart
    data = CategoryChartData()
    data.categories = [p.strftime("%m.%Y") for p in report.trp.index]
    for name in sorted(report.trp.columns):  # порядок серий как в шаблоне — по алфавиту
        data.add_series(name, [float(x) for x in report.trp[name]])
    chart.replace_data(data)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
